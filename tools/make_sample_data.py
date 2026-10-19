"""Generates the ReturnLens sample store (deterministic).

Base files cover orders from 2026-03-16 to 2026-08-31; the "next month" update covers
2026-09-01 to 2026-10-15. Dates are shifted on load so the story always ends "today".

Planted problems (what the agent should find):
  1. EB-200 earbuds, batch B-EB-0712 (supplier SUP-03): charging defects. Inspection INSP-0419 failed.  -> confirmed
  2. MUG-4 mug set: packaging switched to a single-wall box on Jun 15 -> broken/chipped mugs.          -> likely
  3. VASE-1 / LAMP-2 shipped with ParcelPro arrive shattered; carrier damage scans exist.             -> confirmed
  4. SH-LIN linen shirt: new listing photos/description on Jul 20 -> "not as described".              -> likely
  5. SNK-TR sneaker: supplier changed SUP-06 -> SUP-05 on Jun 1 -> runs small; QA doc notes short last. -> confirmed
  6. BLN-1 blender, batch B-BL-0801: lid seal / tamper missing.                                     -> likely
  7. DNM-SL jeans: steady sizing complaints, no batch or supplier pattern.                           -> possible
Update (fixes live Sep 1, so the before/after check has a clear story to tell):
  improved  - MUG-4 back to a double-wall box; EB-200 good batch from Aug 28; SH-LIN listing corrected;
              SNK-TR new batch on a corrected last; BLN-1 bad batch pulled, new batch checks seals.
  not fixed - ParcelPro still breaks glass (the fix "came back" alert); DNM-SL sizing left as is.
"""
from __future__ import annotations

import csv
import random
from datetime import date, timedelta
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "backend" / "app" / "sample_data"
R = random.Random(20261017)
D = lambda s: date.fromisoformat(s)  # noqa: E731

BASE_START, BASE_END = D("2026-03-16"), D("2026-08-31")
UPD_END = D("2026-10-15")

SUPPLIERS = {"SUP-01": "Lumen Manufacturing", "SUP-02": "Kaveri Textiles", "SUP-03": "AudioCore Shenzhen",
             "SUP-04": "ClayWorks Studio", "SUP-05": "Stride Footwear", "SUP-06": "Pace Shoes"}
PRODUCTS = [
    # sku, name, category, supplier, cost, price, weight(share of orders), description
    ("EB-200", "Pulse Earbuds Pro", "Electronics", "SUP-03", 18.0, 59.0, 11,
     "True wireless earbuds with 30-hour battery, USB-C charging case and IPX4 sweat resistance."),
    ("CHG-65", "65W GaN Charger", "Electronics", "SUP-03", 9.0, 29.0, 8,
     "Compact 65W USB-C charger with two ports. Charges a laptop and phone together."),
    ("MUG-4", "Stoneware Mug Set (4)", "Home & Kitchen", "SUP-04", 11.0, 34.0, 10,
     "Set of four 350 ml hand-glazed stoneware mugs. Dishwasher and microwave safe."),
    ("VASE-1", "Fluted Glass Vase", "Home Decor", "SUP-04", 7.0, 26.0, 8,
     "Mouth-blown fluted glass vase, 25 cm tall."),
    ("LAMP-2", "Arc Table Lamp", "Home Decor", "SUP-01", 22.0, 79.0, 6,
     "Brushed brass arc lamp with a frosted glass shade. Bulb included."),
    ("SH-LIN", "Linen Camp Shirt", "Apparel", "SUP-02", 12.0, 45.0, 9,
     "Relaxed camp-collar shirt in 100% linen. Colour: Sage."),
    ("DNM-SL", "Slim Denim Jeans", "Apparel", "SUP-02", 16.0, 59.0, 9,
     "Slim-fit stretch denim. See the size chart for waist and inseam."),
    ("TEE-OR", "Organic Crew Tee", "Apparel", "SUP-02", 5.0, 22.0, 12,
     "Organic cotton crew-neck tee, regular fit."),
    ("SNK-TR", "Trail Runner Sneaker", "Footwear", "SUP-05", 21.0, 85.0, 9,
     "Lightweight trail runner with a grippy outsole. True to size."),
    ("BLN-1", "PowerBlend Blender", "Home & Kitchen", "SUP-01", 28.0, 89.0, 6,
     "1200W blender with 1.8 L jar, lid with seal and tamper included."),
    ("BAG-TT", "Canvas Tote", "Accessories", "SUP-02", 6.0, 24.0, 8,
     "Heavy canvas tote with inner pocket."),
    ("BTL-IN", "Insulated Bottle", "Home & Kitchen", "SUP-01", 7.0, 27.0, 7,
     "750 ml double-wall steel bottle. Keeps drinks cold for 24 hours."),
]
P = {p[0]: p for p in PRODUCTS}

# batches: (batch_id, sku, supplier, received)
BATCHES = []
def batches_for(sku, sup, dates, prefix):
    for d in dates:
        BATCHES.append((f"B-{prefix}-{D(d).strftime('%m%d')}", sku, sup, D(d)))
batches_for("EB-200", "SUP-03", ["2026-03-01", "2026-05-20", "2026-07-12", "2026-08-28"], "EB")
batches_for("CHG-65", "SUP-03", ["2026-03-05", "2026-06-10", "2026-08-25"], "CH")
batches_for("MUG-4", "SUP-04", ["2026-03-02", "2026-05-15", "2026-07-20", "2026-09-10"], "MG")
batches_for("VASE-1", "SUP-04", ["2026-03-04", "2026-06-01", "2026-08-15"], "VS")
batches_for("LAMP-2", "SUP-01", ["2026-03-01", "2026-06-20"], "LP")
batches_for("SH-LIN", "SUP-02", ["2026-03-10", "2026-06-05", "2026-08-20"], "SL")
batches_for("DNM-SL", "SUP-02", ["2026-03-01", "2026-05-25", "2026-08-01"], "DN")
batches_for("TEE-OR", "SUP-02", ["2026-03-01", "2026-05-10", "2026-07-15", "2026-09-01"], "TE")
batches_for("SNK-TR", "SUP-06", ["2026-03-01", "2026-04-25"], "SN")
batches_for("SNK-TR", "SUP-05", ["2026-06-01", "2026-08-05"], "SN")
batches_for("BLN-1", "SUP-01", ["2026-03-12", "2026-06-02", "2026-08-01"], "BL")
batches_for("BAG-TT", "SUP-02", ["2026-03-01", "2026-06-15"], "BG")
batches_for("BTL-IN", "SUP-01", ["2026-03-01", "2026-05-30", "2026-08-20"], "BT")
# fixes that land with the update
batches_for("SNK-TR", "SUP-05", ["2026-08-30"], "SN")   # corrected last (QA-SN-0612 follow-up)
batches_for("BLN-1", "SUP-01", ["2026-08-30"], "BL")    # replaces the quarantined B-BL-0801


def batch_for(sku: str, d: date) -> tuple:
    cands = [b for b in BATCHES if b[1] == sku and b[3] + timedelta(days=2) <= d]
    return max(cands, key=lambda b: b[3])


CARRIERS = [("SwiftShip", 0.45), ("ParcelPro", 0.35), ("EcoFreight", 0.20)]
WAREHOUSES = ["BLR-1", "DEL-2"]
REGIONS = ["Bengaluru", "Delhi", "Mumbai", "Pune", "Hyderabad", "Chennai", "Kolkata"]

TEXT = {
    "defective_eb": ["Left earbud stopped charging after a week", "Case won't charge the right bud anymore",
                     "Right earbud is dead, the charging contacts don't connect", "One earbud keeps dropping out and won't charge",
                     "Stopped working after 10 days, left side has no power", "Charging case lights up but the earbuds never charge"],
    "defective_gen": ["Stopped working after a few days", "Faulty unit, doesn't turn on", "Broke within a week of use"],
    "packaging": ["Box was crushed and two mugs were chipped", "Flimsy box with no dividers, one mug broke",
                  "Packaging too thin, mugs arrived cracked", "No padding at all between the mugs, one shattered",
                  "The box was caved in and a mug handle snapped", "Poorly packed, mugs rattling loose in a thin box"],
    "delivery": ["Arrived shattered, the courier box was smashed", "Glass broke in transit", "Arrived broken, delivery guy dropped it",
                 "Damaged on arrival, shade cracked in shipping", "Came in pieces, outer box had a big dent from the courier",
                 "Shade arrived cracked, parcel was clearly mishandled"],
    "listing": ["Colour is grey, not green like the photos", "Says 100% linen but feels like a polyester blend",
                "Looks nothing like the pictures", "Not as described, the colour is completely different",
                "Fabric is not linen, misleading description", "The photos made it look sage green, it's a dull grey"],
    "sizing_small": ["Runs a full size small", "Too tight in the toe box, had to size up", "Way smaller than my usual size",
                     "Size 9 fits like an 8", "Toes hit the front, order one size up", "Narrow and short compared to the chart"],
    "sizing_jeans": ["Waist runs small", "Too tight around the thighs", "Inseam shorter than the size chart",
                     "Fit is off, too loose at the waist", "Size chart doesn't match the actual fit"],
    "sizing_gen": ["Too big", "Doesn't fit", "Too small"],
    "missing_bln": ["Lid seal was missing from the box", "No tamper included even though the listing says it is",
                    "The rubber gasket for the lid wasn't in the box", "Missing the lid seal, jar leaks",
                    "Box opened, tamper and seal missing"],
    "changed": ["Changed my mind", "No longer needed", "Ordered by mistake", "Found it cheaper elsewhere", "Gift didn't work out"],
    "late": ["Arrived too late for the event", "Took 12 days to deliver", "Delivery was very late"],
    "wrong": ["Received the wrong colour", "Wrong item in the box", "Got a different size than I ordered"],
    "quality": ["Cheap feel, stitching coming loose", "Material feels thin and flimsy", "Colour faded after one wash"],
}
REASON = {"defective_eb": "Defective", "defective_gen": "Defective", "packaging": "Damaged", "delivery": "Damaged on arrival",
          "listing": "Not as described", "sizing_small": "Too small", "sizing_jeans": "Doesn't fit", "sizing_gen": "Doesn't fit",
          "missing_bln": "Missing parts", "changed": "Changed mind", "late": "Arrived late", "wrong": "Wrong item",
          "quality": "Quality not as expected"}


def problem_probs(sku: str, d: date, batch: str, carrier: str) -> dict[str, float]:
    """Per-order probability of each return type."""
    p = {"changed": 0.025, "late": 0.004 + (0.006 if carrier == "EcoFreight" else 0), "wrong": 0.004, "quality": 0.006,
         "defective_gen": 0.005}
    if sku == "EB-200":
        p["defective_eb"] = 0.16 if batch == "B-EB-0712" else 0.018
        p.pop("defective_gen")
    if sku == "MUG-4":
        bad_pack = D("2026-06-15") <= d < D("2026-09-01")
        p["packaging"] = 0.10 if bad_pack else 0.004 if d >= D("2026-09-01") else 0.015
    if sku in ("VASE-1", "LAMP-2"):
        p["delivery"] = (0.11 if sku == "VASE-1" else 0.08) if carrier == "ParcelPro" else 0.012
    if sku == "MUG-4":
        p["delivery"] = 0.03 if carrier == "ParcelPro" else 0.008
    fixed = d >= D("2026-09-01")
    if sku == "SH-LIN":
        p["listing"] = 0.10 if D("2026-07-20") <= d and not fixed else 0.004 if fixed else 0.015
        p["sizing_gen"] = 0.02
    if sku == "SNK-TR":
        p["sizing_small"] = 0.035 if batch == "B-SN-0830" else 0.15 if d >= D("2026-06-03") else 0.035
    if sku == "DNM-SL":
        p["sizing_jeans"] = 0.075
    if sku == "TEE-OR":
        p["sizing_gen"] = 0.02
    if sku == "BLN-1":
        p["missing_bln"] = 0.12 if batch == "B-BL-0801" else 0.008
    return p


def pick_carrier(sku: str) -> str:
    x = R.random()
    acc = 0.0
    for c, w in CARRIERS:
        acc += w
        if x <= acc:
            return c
    return CARRIERS[-1][0]


def build():
    weights = [p[6] for p in PRODUCTS]
    orders, shipments, returns, reviews = [], [], [], []
    oid, rid, vid = 100000, 5000, 9000
    d = BASE_START
    while d <= UPD_END:
        n = int(R.gauss(38, 6) * (1.15 if d.weekday() >= 5 else 1.0) * (1.0 + 0.002 * (d - BASE_START).days))
        for _ in range(max(10, n)):
            oid += 1
            order_id = f"ORD-{oid}"
            lines = R.choices([1, 2], weights=[88, 12])[0]
            skus = R.sample(PRODUCTS, lines) if lines > 1 else [R.choices(PRODUCTS, weights=weights)[0]]
            carrier = pick_carrier(skus[0][0])
            wh = R.choice(WAREHOUSES)
            ship = d + timedelta(days=R.choice([0, 1, 1, 2]))
            deliver = ship + timedelta(days=R.choice([2, 3, 3, 4, 5]) + (3 if carrier == "EcoFreight" and R.random() < 0.3 else 0))
            ship_id = f"SHP-{oid}"
            damaged_order = False
            region = R.choice(REGIONS)
            cust = f"C-{R.randint(10000, 49999)}"
            for prod in skus:
                sku = prod[0]
                qty = 1 if R.random() < 0.9 else 2
                b = batch_for(sku, d)
                # some orders don't carry a batch id (to show inference)
                batch_cell = b[0] if R.random() > 0.08 else ""
                orders.append({"Order ID": order_id, "Order Date": d.isoformat(), "SKU": sku, "Quantity": qty,
                               "Unit Price": f"{prod[5]:.2f}", "Customer ID": cust, "Batch": batch_cell,
                               "Shipment ID": ship_id, "Region": region})
                probs = problem_probs(sku, d, b[0], carrier)
                kind = None
                for k, pr in probs.items():
                    if R.random() < pr:
                        kind = k
                        break
                if kind:
                    if kind == "delivery" and R.random() < 0.75:
                        damaged_order = True
                    lag = R.randint(2, 14) if kind not in ("delivery", "packaging", "wrong", "missing_bln") else R.randint(0, 4)
                    rdate = deliver + timedelta(days=lag)
                    rid += 1
                    has_comment = R.random() < 0.85
                    comment = R.choice(TEXT[kind]) if has_comment else ""
                    reason = REASON[kind]
                    if R.random() < 0.08:
                        reason = "Other"
                    resolution = "Replacement" if kind in ("defective_eb", "missing_bln", "delivery", "packaging") and R.random() < 0.45 else "Refund"
                    rqty = qty if kind not in ("packaging",) else qty
                    refund = round(prod[5] * rqty, 2) if resolution == "Refund" else 0.0
                    label = f"{R.choice([6.5, 7.0, 7.5, 8.0, 9.0]):.2f}" if R.random() < 0.6 else ""
                    cond = ("Damaged" if kind in ("delivery", "packaging") else "Defective" if kind.startswith("defective")
                            else "Incomplete" if kind == "missing_bln" else R.choice(["Unopened", "Opened", "Opened"]))
                    returns.append({"RMA #": f"RMA-{rid}", "Order Number": order_id, "Item SKU": sku,
                                    "Return Date": rdate.isoformat(), "Qty Returned": rqty, "Return Reason": reason,
                                    "Customer Comment": comment, "Refund Amount": f"{refund:.2f}",
                                    "Return Label Cost": label, "Resolution": resolution, "Item Condition": cond})
                    if R.random() < 0.35 and comment:
                        vid += 1
                        reviews.append({"Review ID": f"REV-{vid}", "SKU": sku, "Order ID": order_id,
                                        "Review Date": (rdate - timedelta(days=R.randint(0, 2))).isoformat(),
                                        "Rating": R.choice([1, 1, 2]), "Review": comment})
                elif R.random() < 0.07:
                    vid += 1
                    # happy reviews and a few unhappy ones that didn't return
                    text, rating = R.choice([("Love it, exactly as pictured", 5), ("Great quality for the price", 5),
                                             ("Works well, fast delivery", 4), ("Good, but a bit pricey", 4)])
                    if sku == "SNK-TR" and d >= D("2026-06-03") and b[0] != "B-SN-0830" and R.random() < 0.4:
                        text, rating = R.choice(TEXT["sizing_small"]) + ", kept them anyway", 3
                    if sku == "SH-LIN" and D("2026-07-20") <= d < D("2026-09-01") and R.random() < 0.4:
                        text, rating = "Nice shirt but the colour isn't what the photos show", 3
                    reviews.append({"Review ID": f"REV-{vid}", "SKU": sku, "Order ID": order_id,
                                    "Review Date": (deliver + timedelta(days=R.randint(2, 10))).isoformat(),
                                    "Rating": rating, "Review": text})
            shipments.append({"Tracking ID": ship_id, "Order ID": order_id, "Carrier": carrier, "Warehouse": wh,
                              "Shipped On": ship.isoformat(), "Delivered On": deliver.isoformat(),
                              "Damage Reported": "Yes" if damaged_order else "No"})
        d += timedelta(days=1)
    return orders, shipments, returns, reviews


def write(path: Path, rows: list[dict]):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


def main():
    orders, shipments, returns, reviews = build()
    cut = BASE_END.isoformat()
    split = lambda rows, key: ([r for r in rows if r[key] <= cut], [r for r in rows if r[key] > cut])  # noqa: E731
    o_b, o_u = split(orders, "Order Date")
    s_b, s_u = split(shipments, "Shipped On")
    r_b, r_u = split(returns, "Return Date")
    v_b, v_u = split(reviews, "Review Date")
    r_u = [r for r in r_u if r["Return Date"] <= "2026-10-16"]
    v_u = [r for r in v_u if r["Review Date"] <= "2026-10-16"]
    write(OUT / "orders.csv", o_b)
    write(OUT / "shipments.csv", s_b)
    write(OUT / "returns.csv", r_b)
    write(OUT / "reviews.csv", v_b)
    write(OUT / "update" / "orders_sep_oct.csv", o_u)
    write(OUT / "update" / "shipments_sep_oct.csv", s_u)
    write(OUT / "update" / "returns_sep_oct.csv", r_u)
    write(OUT / "update" / "reviews_sep_oct.csv", v_u)
    write(OUT / "products.csv", [{"SKU": p[0], "Product Name": p[1], "Category": p[2], "Supplier ID": p[3],
                                  "Supplier Name": SUPPLIERS[p[3]], "Unit Cost": f"{p[4]:.2f}", "Price": f"{p[5]:.2f}",
                                  "Description": p[7]} for p in PRODUCTS])
    write(OUT / "batches.csv", [{"Batch ID": b[0], "SKU": b[1], "Supplier ID": b[2], "Supplier Name": SUPPLIERS[b[2]],
                                 "Received Date": b[3].isoformat(), "Units": R.choice([600, 800, 1000, 1200])}
                                for b in BATCHES])
    insp = []
    n = 400
    for b in BATCHES:
        if b[3] > BASE_END:
            continue
        n += 1
        row = {"Inspection ID": f"INSP-{n}", "Batch ID": b[0], "SKU": b[1], "Supplier ID": b[2],
               "Inspection Date": (b[3] + timedelta(days=3)).isoformat(), "Result": "Pass",
               "Defect Rate": f"{R.uniform(0.3, 1.4):.1f}%", "Findings": "Within tolerance"}
        if b[0] == "B-EB-0712":
            row.update({"Inspection ID": "INSP-0419", "Result": "Fail", "Defect Rate": "7.8%",
                        "Findings": "Charging contacts misaligned on left earbud in 7.8% of sampled units; "
                                    "batch released before re-inspection due to stock pressure"})
        if b[0] in ("B-BL-0801", "B-SN-0601", "B-SN-0805"):
            continue  # no routine inspection for these (blender: missing data; sneaker: separate QA report)
        insp.append(row)
    write(OUT / "inspections.csv", insp)
    write(OUT / "change_log.csv", [
        {"Date": "2026-05-10", "SKU": "DNM-SL", "Change Type": "Size chart update", "Supplier ID": "",
         "Description": "Added inseam measurements to the size chart"},
        {"Date": "2026-06-01", "SKU": "SNK-TR", "Change Type": "Supplier change", "Supplier ID": "SUP-05",
         "Description": "Moved production from Pace Shoes (SUP-06) to Stride Footwear (SUP-05)"},
        {"Date": "2026-06-15", "SKU": "MUG-4", "Change Type": "Packaging change", "Supplier ID": "",
         "Description": "Switched to a single-wall mailer box without dividers to cut shipping weight"},
        {"Date": "2026-07-20", "SKU": "SH-LIN", "Change Type": "Listing update", "Supplier ID": "",
         "Description": "New product photos and description: colour 'Sage', fabric '100% linen'"},
        {"Date": "2026-08-12", "SKU": "LAMP-2", "Change Type": "Listing update", "Supplier ID": "",
         "Description": "Added a lifestyle photo"},
    ])
    write(OUT / "update" / "change_log_sep_oct.csv", [
        {"Date": "2026-09-01", "SKU": "MUG-4", "Change Type": "Packaging change", "Supplier ID": "",
         "Description": "Back to a double-wall box with cardboard dividers"},
        {"Date": "2026-09-01", "SKU": "SH-LIN", "Change Type": "Listing update", "Supplier ID": "",
         "Description": "Re-shot photos in daylight, colour renamed 'Grey sage', fabric corrected to 'linen blend'"},
        {"Date": "2026-08-30", "SKU": "SNK-TR", "Change Type": "Other", "Supplier ID": "SUP-05",
         "Description": "Stride Footwear ships batch B-SN-0830 on the corrected last (QA-SN-0612 corrective action)"},
        {"Date": "2026-08-30", "SKU": "BLN-1", "Change Type": "Other", "Supplier ID": "SUP-01",
         "Description": "Batch B-BL-0801 pulled from sale; new batch B-BL-0830 with a packing check for lid seal and tamper"},
    ])
    from docx import Document
    doc = Document()
    doc.add_heading("Quality Inspection Report", 1)
    for line in ["Report ID: QA-SN-0612", "Batch: B-SN-0601", "SKU: SNK-TR", "Supplier: SUP-05 (Stride Footwear)",
                 "Inspection date: 2026-06-12", "Result: Pass with deviations", "Defect rate: 2.1%",
                 "Findings: Last length measured 4-6 mm shorter than the size spec on sizes 8 to 10. Toe box width within "
                 "tolerance. Approved for sale pending supplier corrective action on the last."]:
        doc.add_paragraph(line)
    doc.add_paragraph("Inspector: R. Menon, Stride Footwear QA")
    doc.save(OUT / "QA_report_B-SN-0601.docx")
    print(f"orders {len(o_b)}+{len(o_u)}  returns {len(r_b)}+{len(r_u)}  reviews {len(v_b)}+{len(v_u)}  shipments {len(s_b)}+{len(s_u)}")


if __name__ == "__main__":
    main()
